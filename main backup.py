import sys
import math
import sqlite3
from datetime import datetime
from PySide6.QtWidgets import (
    QCheckBox, QDialog, QVBoxLayout, QDialogButtonBox,
    QApplication, QMainWindow, QFileDialog, QMessageBox, QInputDialog,
    QPushButton, QHBoxLayout, QWidget, QLabel, QListWidgetItem
)
from PySide6.QtCore import Qt, QPropertyAnimation, QParallelAnimationGroup, QEvent, QTimer
from ui_interface import Ui_MainWindow
import fitz
from docx import Document
import textwrap
from docx.shared import Inches, Pt  
from database import (
    init_db, save_meeting, get_all_recordings,
    get_meeting_by_id, delete_meeting_by_id
)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)

        # State variables
        self.drawer_open = True
        self.is_maximized = False
        self.attendees = []
        self.agendas = []            
        self.speaker_map = {}
        self.current_meeting_id = None
        self.current_draft_id = None
        self.current_log_title = None
        self.current_log_agenda = None
        self.raw_transcript_text = ""  # Stores actual conversation text
        self.transcript_html_template = ""  # Stores base HTML structure

        # Initialize UI & DB
        self._init_window()
        self._connect_signals()
        init_db()
        self._load_saved_recordings()
        self.show()

        # Initial button states
        self.ui.proceedBtn.setEnabled(True)
        self.ui.saveBtn.setEnabled(True)
        self._update_button_states()  

    def _init_window(self):
        self.setFixedSize(1000, 600)
        self.setWindowFlags(Qt.FramelessWindowHint)
        self.ui.sideMenuContainer.setMinimumWidth(200)
        self.ui.sideMenuContainer.setMaximumWidth(200)
        

        # Set tooltip for the clear button
        self.ui.clearBtn.setToolTip("Clear Inputs")
        
        # Always show these controls
        self.ui.assignSpeakerBtn.setVisible(True)
        self.ui.replaceTextBtn.setVisible(True)
        self.ui.speakerNumber.setVisible(True)
        self.ui.speakerName.setVisible(True)
        
        # Initially disable inputs
        self.ui.speakerNumber.setEnabled(False)
        self.ui.speakerName.setEnabled(False)
        self.ui.assignSpeakerBtn.setEnabled(False)

    def _connect_signals(self):
        u = self.ui
        # Navigation
        u.homeButton.clicked.connect(lambda: self._set_page(u.HomePage))
        u.returnBtn.clicked.connect(lambda: self._set_page(u.GeneratePage1))
        u.LogButton.clicked.connect(lambda: self._set_page(u.activityListPage))
        u.returnLogBtn.clicked.connect(lambda: self._set_page(u.activityListPage))
        u.getStartedBtn.clicked.connect(lambda: self._set_page(u.GeneratePage1))
        u.clearBtn.clicked.connect(self._clear_and_go_home)
        u.logSaveBtn.clicked.connect(self._save_log_to_database)
        # File & export
        u.uploadBtn.clicked.connect(self._open_file_manager)
        u.proceedBtn.clicked.connect(self._save_transcript_and_proceed)
        # Transcription and generation buttons
        u.transcribeBtn.clicked.connect(self._transcribe_audio)
        u.generateBtn.clicked.connect(self._generate_minutes)
        u.logGenerateBtn.clicked.connect(self._generate_log_minutes)
        # Window controls
        u.close.clicked.connect(self.close)
        u.minimizeWindow.clicked.connect(self.showMinimized)
        u.maximizeWindow.clicked.connect(self._toggle_maximize)
        u.sideDrawerButton.clicked.connect(self._toggle_drawer)
        # Text inputs
        u.attendeesInput.installEventFilter(self)
        u.agendaInput.installEventFilter(self)
        u.assignSpeakerBtn.clicked.connect(self._assign_speaker_name)
        u.replaceTextBtn.clicked.connect(self._replace_in_minutes)
        u.attendeesInput.textChanged.connect(self._update_attendees_list)
        u.agendaInput.textChanged.connect(self._update_agendas_list)
        # Database ops
        u.saveBtn.clicked.connect(self._save_to_database)
        u.savedListWidget.itemClicked.connect(self._load_selected_log)
        u.transcriptDisplay.textChanged.connect(self._update_button_states)
        u.minutesDisplay.textChanged.connect(self._update_button_states)

    def _update_attendees_list(self):
        """Update attendees list from comma-separated input"""
        raw_text = self.ui.attendeesInput.toPlainText().strip()
        self.attendees = [name.strip() for name in raw_text.split(',') if name.strip()]
        self._update_transcript()

    def _update_agendas_list(self):
        """Update agenda list from comma-separated input"""
        raw_text = self.ui.agendaInput.toPlainText().strip()
        self.agendas = [agenda.strip() for agenda in raw_text.split(',') if agenda.strip()]
        self._update_transcript()

    def _transcribe_audio(self):
            #ITO SAYO JOE, SAMPLE LANG TO
            file_path = self.ui.uploadedFileLabel.text()
            if not file_path:
                QMessageBox.warning(self, "Error", "Please upload an audio file first.")
                return
            
            # Store current scroll position, KEEP THIS DONT DELETE PLS! THANK YOU!
            scroll = self.ui.transcriptDisplay.verticalScrollBar().value()
            
            self.ui.transcriptDisplay.setPlainText("Transcribing... Please wait.")
            QApplication.processEvents()
            
            QTimer.singleShot(1500, lambda: self._complete_transcription(scroll))
    
    #ITO NAKA FORMA NA, STORE MO SA transcript VARIABLE MAGIGING OUTCOME NG TRANSCRIPTION
    def _complete_transcription(self, original_scroll):
        transcript = self._generate_sample_transcript()
        current_time = datetime.now().strftime("%B %d, %Y | %I:%M %p")
        agenda_str = self._format_agendas()
        attendees_str = self._format_attendees(self.attendees)
        
        # Store raw transcript and template separately
        self.raw_transcript_text = transcript
        self.transcript_html_template = f"""
        <b>Meeting Date:</b> {current_time}<br><br>
        <b>Meeting Agenda:</b><br><pre>{{AGENDA}}</pre><br><br>
        <b>Meeting Attendees:</b><br><pre>{{ATTENDEES}}</pre><br>
        <b>Meeting Transcript:</b><br><pre>{{TRANSCRIPT}}</pre>
        """
        
        # Initial render
        html = self.transcript_html_template.format(
            AGENDA=agenda_str,
            ATTENDEES=attendees_str,
            TRANSCRIPT=self.raw_transcript_text
        )
        self.ui.transcriptDisplay.setHtml(html)
        self.ui.transcriptDisplay.verticalScrollBar().setValue(original_scroll)
        self._update_button_states()
        
    def _update_button_states(self):
        #Updates Buttons and Condition, PLEASE KEEP!
        has_transcript = bool(self.ui.transcriptDisplay.toPlainText().strip())
        has_minutes = bool(self.ui.minutesDisplay.toPlainText().strip())
    
    # Enable speaker controls only if transcript exists. PLEASE KEEP!
        self.ui.speakerNumber.setEnabled(has_transcript)
        self.ui.speakerName.setEnabled(has_transcript)
        self.ui.assignSpeakerBtn.setEnabled(has_transcript)
        
        # Show/hide other buttons. PLEASE KEEP!
        self.ui.replaceTextBtn.setVisible(has_minutes)

    def _generate_sample_transcript(self):
        #SAMPLE LANG TO, YOU CAN DELETE THIS JOE IF MERON NG MODEL
        return """
Speaker 001: Good morning everyone, thank you for joining today's meeting.
Speaker 002: Good morning. I have the quarterly reports ready to discuss.
Speaker 001: Great. Before we start, let's quickly go through the agenda.
Speaker 003: I also have some updates regarding the new project timeline.
Speaker 001: Perfect. Let's start with the first agenda item then.
Speaker 002: The first quarter results show a 15% increase in revenue.
Speaker 003: That's better than we projected in our last meeting.
Speaker 001: Indeed. Let's discuss how we can maintain this momentum.
"""

    def _generate_minutes(self):
        #DITO KA MAGCODE LEX, SAMPLE LAG TONG NANDITO, YOU CAN ENTIRELY DELETE THIS AND CREATE YOUR OWN PROCESS
        transcript = self.ui.transcriptDisplay.toPlainText()
        
        if not transcript or "Meeting Transcript:" not in transcript:
            QMessageBox.warning(self, "Error", "No transcript available. Please transcribe an audio file first.")
            return
            
        self.ui.minutesDisplay.setPlainText("Generating minutes... Please wait.")
        QApplication.processEvents()
        
        QTimer.singleShot(1500, lambda: self._complete_minutes_generation())
            
    def _complete_minutes_generation(self):
        transcript_text = self.ui.transcriptDisplay.toPlainText()
        date_line = ""
        for line in transcript_text.split('\n'):
            if "Meeting Date:" in line:
                date_line = line
                break

#SAMPLE FORMATTING OF CONTENTS
        minutes = f"""

MINUTES OF THE MEETING
{date_line}

ATTENDEES:
{', '.join(self.attendees) if self.attendees else "No attendees recorded"}

AGENDA:
{self._format_agendas_for_minutes()}

DISCUSSION POINTS:
1. The team discussed quarterly results, which showed a 15% increase in revenue.
2. Project timeline updates were presented and reviewed.
3. Discussion on strategies to maintain growth momentum.

ACTION ITEMS:
1. Team members to review detailed quarterly reports by next week.
2. Update project timeline documentation with new milestones.
3. Schedule follow-up meeting to discuss implementation strategies.

NEXT MEETING:
To be scheduled for next month.
"""
        self.ui.minutesDisplay.setPlainText(minutes)
        self._update_button_visibility()


    def _format_agendas_for_minutes(self):
        if not self.agendas:
            return "No agenda items recorded"
        return "\n".join([f"{i+1}. {agenda}" for i, agenda in enumerate(self.agendas)])

    def _generate_log_minutes(self):
    ##SAME FUNCTIONS HERE PERO SA ACTIVITY LOG GENERATE BUTTON TO, YOU CAN ALSO RECREATE THIS PROCESS HERE
        transcript = self.ui.logTranscriptDisplay.toPlainText()
        
        if not transcript:
            QMessageBox.warning(self, "Error", "No transcript available in the log.")
            return
            
        self.ui.logMinutesDisplay.setPlainText("Generating minutes... Please wait.")
        QApplication.processEvents()
        
        QTimer.singleShot(1500, lambda: self._complete_log_minutes_generation())

    def _complete_log_minutes_generation(self):
        transcript = self.ui.logTranscriptDisplay.toPlainText()
        date_match = None
        attendees = []
        agenda_items = []
        
        lines = transcript.split('\n')
        for line in lines:
            if "Meeting Date:" in line:
                date_match = line.strip()
            elif line.startswith("•"):
                agenda_items.append(line.strip()[2:].strip())
                
        minutes = f"""
MINUTES OF THE MEETING
{date_match if date_match else datetime.now().strftime("%B %d, %Y | %I:%M %p")}

AGENDA:
{self._format_list_items(agenda_items)}

DISCUSSION SUMMARY:
1. The team reviewed current progress on ongoing projects.
2. Key challenges were identified and potential solutions were discussed.
3. New opportunities for growth were explored.

ACTION ITEMS:
1. Team members to complete assigned tasks before the next meeting.
2. Follow-up meetings to be scheduled with specific stakeholders.
3. Progress report to be prepared for management review.

NEXT STEPS:
The team will reconvene next week to assess progress on action items.
"""
        self.ui.logMinutesDisplay.setPlainText(minutes)


    def _format_list_items(self, items):
        if not items:
            return "No items recorded"
        return "\n".join([f"{i+1}. {item}" for i, item in enumerate(items)])

    def eventFilter(self, obj, event):
        # Handle Enter key for comma separation
        if (obj in (self.ui.attendeesInput, self.ui.agendaInput) and 
            event.type() == QEvent.KeyPress and 
            event.key() == Qt.Key_Return):
            
            cursor = obj.textCursor()
            cursor.insertText(", ")
            return True
            
        return super().eventFilter(obj, event)

    def _format_attendees(self, names, cols=3):
        if not names:
            return "No data input"
        rows = math.ceil(len(names) / cols)
        lines = []
        for r in range(rows):
            row = []
            for c in range(cols):
                i = r + c * rows
                if i < len(names):
                    row.append(f"{i+1:>3}. {names[i]:<25}")
            lines.append("".join(row).rstrip())
        return "\n".join(lines)

    def _format_agendas(self):
        if not self.agendas:
            return "No data input"
        return "\n".join([f"• {agenda}" for agenda in self.agendas])

    def _update_transcript(self):
        agenda_str = self._format_agendas()
        attendees_str = self._format_attendees(self.attendees)
        filename = self.ui.uploadedFileLabel.text()
        
        # Show filename if no transcript exists
        transcript_content = ""
        if not self.raw_transcript_text:
            if filename:
                transcript_content = f"[Uploaded File: {filename}]"
        else:
            transcript_content = self.raw_transcript_text
        
        # Rebuild HTML using template
        html = self.transcript_html_template.format(
            AGENDA=agenda_str,
            ATTENDEES=attendees_str,
            TRANSCRIPT=transcript_content
        )
        self.ui.transcriptDisplay.setHtml(html)
        self._update_button_visibility()

    def _open_file_manager(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Open File", "",
            "Audio/Video Files (*.mp3 *.mp4 *.wav)"
        )
        if path:
            self.ui.uploadedFileLabel.setText(path.split("/")[-1])

    def _wrap_text(self, text, width):
        wrapped = []
        for line in text.split('\n'):
            wrapped_lines = textwrap.wrap(line, width=width, break_long_words=True)
            wrapped.extend(wrapped_lines or [''])
        return wrapped

    def _save_pdf(self, file_path, sections):
        doc = fitz.open()
        w, h = 595, 842
        margin = 50
        line_height = 14

        for title, text in sections:
            page = doc.new_page(width=w, height=h)
            y = margin
            page.insert_text((margin, y), title,
                             fontsize=14, fontname="helv", render_mode=3)
            y += 2 * line_height
            
            wrapped_lines = self._wrap_text(text, 70)
            font_size = 9
            
            for line in wrapped_lines:
                if y + line_height > h - margin:
                    page = doc.new_page(width=w, height=h)
                    y = margin
                page.insert_text((margin, y), line,
                                 fontsize=font_size, fontname="courier")
                y += line_height
        doc.save(file_path)

    def _save_word(self, file_path, sections):
        doc = Document()
        for title, text in sections:
            p = doc.add_paragraph()
            r = p.add_run(title)
            r.bold = True
            wrapped_lines = self._wrap_text(text, 80)
            
            for line in wrapped_lines:
                if line.startswith(("Meeting Agenda", "Meeting Attendees")):
                    p = doc.add_paragraph()
                    r = p.add_run(line)
                    r.bold = True
                else:
                    doc.add_paragraph(line)
        doc.save(file_path)


    def _assign_speaker_name(self):
        # Get current scroll position
        scroll = self.ui.transcriptDisplay.verticalScrollBar().value()
        
        if not self.raw_transcript_text:  # Changed condition
            QMessageBox.warning(self, "Error", "Please transcribe an audio file first.")
            return

        speaker_number = self.ui.speakerNumber.toPlainText().strip()
        speaker_name = self.ui.speakerName.toPlainText().strip()
        
        # Validation
        if not speaker_number:
            QMessageBox.warning(self, "Missing Information", "Please enter a speaker number.")
            return
            
        if not speaker_name:
            QMessageBox.warning(self, "Missing Information", "Please enter a speaker name.")
            return
            
        speaker_tag = f"Speaker {speaker_number}" if not speaker_number.startswith("Speaker ") else speaker_number
            
        # Operate on raw_transcript_text instead of display text
        original_text = self.raw_transcript_text  # Changed source
        new_text = original_text.replace(speaker_tag, speaker_name)
        
        if original_text == new_text:
            QMessageBox.warning(self, "Not Found", 
                            f"Speaker tag '{speaker_tag}' not found in transcript.")
            return
            
        # Update both the stored raw text and display
        self.raw_transcript_text = new_text  # Store modified text
        self.speaker_map[speaker_tag] = speaker_name
        
        # Refresh display using the template
        self._update_transcript()  # This will rebuild the HTML
        
        # Restore scroll position
        self.ui.transcriptDisplay.verticalScrollBar().setValue(scroll)
        
        self.ui.speakerNumber.clear()
        self.ui.speakerName.clear()

    def _update_transcript_with_speaker_names(self):
        raw = self.ui.transcriptDisplay.toPlainText()
        for tag, name in self.speaker_map.items():
            raw = raw.replace(tag, name)
        self.ui.transcriptDisplay.setPlainText(raw)
        self._update_button_visibility()

    def _replace_in_minutes(self):
        find_txt, ok1 = QInputDialog.getText(
            self, "Find Text", "Enter text to find:"
        )
        if not ok1:
            return
        repl_txt, ok2 = QInputDialog.getText(
            self, "Find Text", "Enter text to replace with:"
        )
        if not ok2:
            return
        orig = self.ui.minutesDisplay.toPlainText()
        self.ui.minutesDisplay.setPlainText(orig.replace(find_txt, repl_txt))
        self._update_button_visibility()

    def _update_button_visibility(self):
        has_transcript = bool(self.ui.transcriptDisplay.toPlainText().strip())
        has_minutes    = bool(self.ui.minutesDisplay.toPlainText().strip())
        self.ui.assignSpeakerBtn.setVisible(has_transcript)
        self.ui.replaceTextBtn.setVisible(has_minutes)

    def _toggle_drawer(self):
        start = self.ui.sideMenuContainer.width()
        end = 0 if start > 0 else 200
        self.drawer_open = (end > 0)
        self._animate_drawer(start, end)

    def _animate_drawer(self, start_width, end_width):
        if getattr(self, "_drawer_anim", None) and self._drawer_anim.state():
            self._drawer_anim.stop()

        group = QParallelAnimationGroup(self)
        for prop in (b"minimumWidth", b"maximumWidth"):
            anim = QPropertyAnimation(self.ui.sideMenuContainer, prop)
            anim.setDuration(250)
            anim.setStartValue(start_width)
            anim.setEndValue(end_width)
            group.addAnimation(anim)
        group.start()
    def _clear_and_go_home(self):
        """
        Clear all inputs and content from GeneratePage1 and GeneratePage2 and navigate back to the HomePage
        """
        # Ask the user to confirm
        confirm = QMessageBox.question(
            self, "Confirm Clear",
            "Are you sure you want to clear all content and return to the home page?\n"
            "Any unsaved work will be lost.",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if confirm == QMessageBox.Yes:
            # Clear all inputs and content
            self._clear_all_content()
            
            # Navigate back to the home page
            self._set_page(self.ui.HomePage)
            

    def _clear_all_content(self):
        """
        Clear all inputs and content from GeneratePage1 and GeneratePage2
        """
        # Clear file upload
        self.ui.uploadedFileLabel.setText("")
        
        # Clear attendees
        self.ui.attendeesInput.clear()
        self.attendees.clear()
        
        
        # Clear agenda items
        self.ui.agendaInput.clear()
        self.agendas.clear()
        
        # Clear transcript and minutes displays
        self.ui.transcriptDisplay.clear()
        self.ui.minutesDisplay.clear()
        self.raw_transcript_text = ""
        self.transcript_html_template = ""
        
        # Clear speaker inputs
        self.ui.speakerNumber.clear()
        self.ui.speakerName.clear()
        
        # Reset speaker mappings
        self.speaker_map.clear()
        
        # Reset meeting IDs
        self.current_meeting_id = None
        self.current_draft_id = None
        
        # Reset UI elements
        self._update_button_states()
        self._update_button_visibility()
        
        # Reset other displays if needed
        self.ui.logTranscriptDisplay.clear()
        self.ui.logMinutesDisplay.clear()
        
    def _toggle_maximize(self):
        if self.is_maximized:
            self.showNormal()
        else:
            self.showMaximized()
        self.is_maximized = not self.is_maximized

    def _save_transcript_and_proceed(self):
        raw_text = self.ui.transcriptDisplay.toPlainText()
        lines = raw_text.splitlines()
        
        content_lines = [
            line for line in lines 
            if line.strip() and 
            not line.startswith(("Meeting Date:", "Meeting Agenda:", "Meeting Attendees:", "Meeting Transcript:"))
        ]

        is_empty = (
            not self.agendas and 
            not self.attendees and 
            not any("•" in line for line in lines) and
            not any(line.strip().isdigit() for line in lines) and
            not content_lines
        )

        if is_empty:
            QMessageBox.warning(
                self,
                "Empty Draft",
                "Cannot save empty draft!\nPlease transcribe a recording or add meeting details first."
            )
            return

        draft_title = f"Transcript Draft {datetime.now().strftime('%m/%d/%Y at %I:%M %p')}"
        agenda_text = "\n".join(self.agendas).strip()
        transcript = self.ui.transcriptDisplay.toPlainText().strip()
        minutes = ""

        draft_id = save_meeting(
            draft_title, agenda_text, transcript, minutes,
            is_draft=1, meeting_id=self.current_draft_id
        )
        
        if draft_id:
            self.current_draft_id = draft_id
            QMessageBox.information(self, "Saved", "Draft updated successfully.")
            self._load_saved_recordings()
            self.ui.proceedBtn.setEnabled(False)
            QTimer.singleShot(1000, lambda: self.ui.proceedBtn.setEnabled(True))
            self._set_page(self.ui.GeneratePage2)
        else:
            QMessageBox.warning(self, "Error", "Failed to save draft.")

    def _save_to_database(self):
        # Create export dialog
        dialog = QDialog(self)
        dialog.setWindowTitle("Click Ok to Save Record")
        layout = QVBoxLayout()
        
        # Create checkboxes
        transcript_cb = QCheckBox("Export Transcript")
        minutes_cb = QCheckBox("Export Minutes")
        
        # Create button box
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        
        # Add widgets to layout
        layout.addWidget(transcript_cb)
        layout.addWidget(minutes_cb)
        layout.addWidget(button_box)
        dialog.setLayout(layout)
        
        # Connect buttons
        button_box.accepted.connect(dialog.accept)
        button_box.rejected.connect(dialog.reject)
        
        # Show dialog
        if dialog.exec() != QDialog.Accepted:
            return
        
        # Get user choices
        export_transcript = transcript_cb.isChecked()
        export_minutes = minutes_cb.isChecked()
        
        # Handle export
        sections = []
        text = self.ui.transcriptDisplay.toPlainText().strip()
        mins = self.ui.minutesDisplay.toPlainText().strip()
        
        if export_transcript and text:
            sections.append(("Meeting Transcript", text))
        if export_minutes and mins:
            sections.append(("Minutes of the Meeting", mins))
        
        if sections:
            path, _ = QFileDialog.getSaveFileName(
                self, "Save File", "", 
                "PDF Files (*.pdf);;Word Files (*.docx)"
            )
            if path:
                if path.endswith(".pdf"):
                    self._save_pdf(path, sections)
                elif path.endswith(".docx"):
                    self._save_word(path, sections)
                QMessageBox.information(self, "Exported", f"File saved to:\n{path}")

        # Always save to database
        raw_text = self.ui.transcriptDisplay.toPlainText()
        mins_text = self.ui.minutesDisplay.toPlainText().strip()
        
        is_empty = (
            not self.agendas and 
            not self.attendees and 
            not raw_text.strip() and 
            not mins_text
        )

        if is_empty:
            QMessageBox.warning(
                self,
                "Empty Meeting",
                "Cannot save empty meeting!\nPlease add content before saving."
            )
            return

        final_title = f"Record {datetime.now().strftime('%m/%d/%Y at %I:%M %p')}"
        agenda_text = "\n".join(self.agendas).strip()
        transcript = raw_text.strip()
        minutes = mins_text

        final_id = save_meeting(final_title, agenda_text, transcript, minutes, is_draft=0)
        
        if final_id:
            if self.current_draft_id:
                delete_meeting_by_id(self.current_draft_id)
                self.current_draft_id = None
            
            QMessageBox.information(self, "Saved", "Record Saved successfully.")
            self.ui.saveBtn.setEnabled(False)
            QTimer.singleShot(1000, lambda: self.ui.saveBtn.setEnabled(True))
            self._load_saved_recordings()
            self._clear_generate_pages()
            self._set_page(self.ui.GeneratePage1)  # Navigate back to GeneratePage1
        else:
            QMessageBox.warning(self, "Error", "Failed to save the meeting.")

    def _save_log_to_database(self):
        if not self.current_meeting_id:
            QMessageBox.warning(self, "Error", "No meeting selected!")
            return

        # Create export dialog
        dialog = QDialog(self)
        dialog.setWindowTitle("Click Ok to Save Record")
        layout = QVBoxLayout()
        
        transcript_cb = QCheckBox("Export Transcript")
        minutes_cb = QCheckBox("Export Minutes")
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        
        layout.addWidget(transcript_cb)
        layout.addWidget(minutes_cb)
        layout.addWidget(button_box)
        dialog.setLayout(layout)
        
        button_box.accepted.connect(dialog.accept)
        button_box.rejected.connect(dialog.reject)
        
        if dialog.exec() != QDialog.Accepted:
            return

        # Get user choices
        export_transcript = transcript_cb.isChecked()
        export_minutes = minutes_cb.isChecked()
        
        # Handle export
        sections = []
        transcript = self.ui.logTranscriptDisplay.toPlainText().strip()
        minutes = self.ui.logMinutesDisplay.toPlainText().strip()
        
        if export_transcript and transcript:
            sections.append(("Meeting Transcript", transcript))
        if export_minutes and minutes:
            sections.append(("Meeting Minutes", minutes))
        
        if sections:
            path, _ = QFileDialog.getSaveFileName(
                self, "Save File", "",
                "PDF Files (*.pdf);;Word Files (*.docx)"
            )
            if path:
                if path.endswith(".pdf"):
                    self._save_pdf(path, sections)
                elif path.endswith(".docx"):
                    self._save_word(path, sections)
                QMessageBox.information(self, "Exported", f"File saved to:\n{path}")

        # Check if this is a draft and modify title if needed
        title = self.current_log_title
        if title and title.startswith("Transcript Draft"):
            # Replace "Transcript Draft" with "Recording" but keep the date and time
            title = title.replace("Transcript Draft", "Recording")

        # Update database
        success = save_meeting(
            title=title,  # Use the potentially modified title
            agenda=self.current_log_agenda,
            transcript=transcript,
            minutes=minutes,
            is_draft=0,  # Set is_draft to 0 when converting from draft to recording
            meeting_id=self.current_meeting_id
        )
        
        if success:
            QMessageBox.information(self, "Saved", "Meeting updated successfully!")
            self._load_saved_recordings()  # Refresh list
        else:
            QMessageBox.warning(self, "Error", "Failed to save changes!")

    def _clear_generate_pages(self):
        self.ui.agendaInput.clear()
        self.agendas.clear()
        self.ui.attendeesInput.clear()
        self.attendees.clear()
        self._update_transcript()
        self.ui.transcriptDisplay.clear()
        self.ui.minutesDisplay.clear()
        self.ui.uploadedFileLabel.setText("")
        self.speaker_map.clear()
        self.current_draft_id = None

    def _load_saved_recordings(self):
        self.ui.savedListWidget.clear()
        for meeting_id, title, is_draft in get_all_recordings():
            item = QListWidgetItem()
            self.ui.savedListWidget.addItem(item)

            item_widget = QWidget()
            layout = QHBoxLayout(item_widget)
            layout.setContentsMargins(5, 0, 5, 0)
            
            text_label = QLabel(title)
            text_label.setStyleSheet("""
                QLabel {
                    color: rgb(41, 28, 14);
                    font-family: 'Courier New';
                    font-size: 16px;
                    padding: 3px;
                    background-color: transparent;
                }
            """)
            text_label.setProperty("meeting_id", meeting_id)
            
            layout.addWidget(text_label)

            delete_btn = QPushButton("Delete")
            delete_btn.setFixedWidth(60)
            delete_btn.setStyleSheet("""
                QPushButton {
                    background-color:  rgb(185, 128, 118);
                    color: white;
                    border-radius: 3px;
                    padding: 3px;
                }
                QPushButton:hover {
                    background-color: #e03131;
                }
            """)
            delete_btn.clicked.connect(lambda _, m_id=meeting_id: self._delete_meeting(m_id))

            layout.addStretch(1)
            layout.addWidget(delete_btn)
            self.ui.savedListWidget.setItemWidget(item, item_widget)

    def _delete_meeting(self, meeting_id):
        confirm = QMessageBox.question(
            self, "Confirm Delete",
            f"Are you sure you want to delete this recording?",
            QMessageBox.Yes | QMessageBox.No
        )
        if confirm == QMessageBox.Yes:
            if delete_meeting_by_id(meeting_id):
                QMessageBox.information(self, "Deleted", "Recording deleted successfully.")
                cur = self.ui.stackedWidget.currentWidget()
                if cur == self.ui.activityLogPage and meeting_id == self.current_meeting_id:
                    self._set_page(self.ui.activityListPage)
                self._load_saved_recordings()
            else:
                QMessageBox.warning(self, "Error", "Failed to delete the recording.")

    def _load_selected_log(self, item):
        widget = self.ui.savedListWidget.itemWidget(item)
        if widget:
            label = widget.findChild(QLabel)
            if label:
                meeting_id = label.property("meeting_id")
                if meeting_id:
                    self.current_meeting_id = meeting_id
                    # Get title, agenda, transcript, minutes
                    title, agenda, transcript, minutes = get_meeting_by_id(meeting_id)
                    # Store title and agenda for saving
                    self.current_log_title = title
                    self.current_log_agenda = agenda
                    # Update displays
                    self.ui.logTranscriptDisplay.setPlainText(transcript)
                    self.ui.logMinutesDisplay.setPlainText(minutes)
                    self._set_page(self.ui.activityLogPage)

    def _set_page(self, page):
        self.ui.stackedWidget.setCurrentWidget(page)

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    sys.exit(app.exec())